"""Import the supplied EKP workbook without losing conditional service columns."""
import hashlib
import re
from pathlib import Path

from openpyxl import load_workbook

FLAGS = {
    'victims': 'Пострадавшие / погибшие', 'access_blocked': 'Нет доступа / заблокированы',
    'life_danger': 'Угроза людям', 'offense': 'Правонарушение',
    'victims_away': 'Пострадавшие не на месте / отказ от скорой',
    'gasified': 'Газификация', 'medical_help': 'Медицинская помощь',
    'evacuation': 'Требуется эвакуация', 'mass_event': 'Более 5 человек / опасное действие',
    'traffic_blocked': 'Перекрытие движения', 'tunnel': 'Тоннель',
    'pedestrian_structure': 'Пешеходное сооружение', 'vehicle_structure': 'Автомобильное сооружение',
    'communications_object': 'Объект связи', 'construction': 'Стройка',
    'culture_object': 'Объект культуры из перечня', 'polygon': 'Событие на полигоне',
}
# Excel column numbers. Alternatives come from the original three-row header.
CONDITIONS = {
    14: ('not', ['access_blocked']), 15: ('any', ['access_blocked']),
    16: ('not', ['life_danger', 'victims', 'access_blocked']),
    17: ('any', ['life_danger']), 18: ('any', ['victims']), 19: ('any', ['access_blocked']),
    21: ('not', ['offense', 'victims']), 22: ('any', ['offense']), 23: ('any', ['victims']),
    24: ('not', ['victims', 'victims_away']), 25: ('victims_here', []), 26: ('any', ['victims_away']),
    27: ('not', ['gasified']), 28: ('any', ['gasified']),
    29: ('not', ['life_danger', 'victims', 'medical_help', 'evacuation']),
    30: ('any', ['life_danger']), 31: ('any', ['victims']),
    32: ('any', ['medical_help']), 33: ('any', ['evacuation']),
    34: ('not', ['mass_event']), 35: ('any', ['mass_event']),
    38: ('not', ['victims', 'traffic_blocked']), 39: ('any', ['victims']), 40: ('any', ['traffic_blocked']),
    42: ('not', ['tunnel', 'pedestrian_structure', 'vehicle_structure']),
    43: ('any', ['tunnel']), 44: ('any', ['pedestrian_structure']), 45: ('any', ['vehicle_structure']),
    48: ('any', ['communications_object']), 78: ('not', ['construction']), 79: ('any', ['construction']),
    82: ('any', ['culture_object']), 93: ('not', ['traffic_blocked']), 94: ('any', ['traffic_blocked']),
    96: ('any', ['polygon']), 97: ('not', ['polygon']),
}
SERVICE_NAMES = {
    '101': 'Служба 101', '102': 'Служба 102', '103': 'Служба 103', '104': 'МОСГАЗ / Служба 104',
    'PSC': 'ОДС ПСЦ', 'MGPSS': 'МГПСС', 'CEMP': 'ЦЭМП', 'FSB': 'ФСБ',
    'MOSOBLGAZ': 'Мособлгаз', 'GORHOZ': 'Городское хозяйство', 'GORMOST': 'Гормост',
    'MOSBEZ': 'ГКУ МОСБЕЗ', 'MOSBEZ_ANALYTICS': 'МОСБЕЗ · аналитика',
    'TERRITORY': 'Территориальные ОИВ', 'TINAO': 'Территориальные ОИВ ТиНАО',
}

ACTIVE_CLASSIFIER = 'Классификатор_происшествий_v_046_24_корректировка_МВД_+_Департамент.xlsx'
def active_workbook(root='source_materials'):
    import os
    name = os.getenv('CLASSIFIER_FILE', ACTIVE_CLASSIFIER)
    if Path(name).name != name: raise ValueError('CLASSIFIER_FILE должен содержать только имя файла')
    path = Path(root) / name
    if not path.is_file(): raise ValueError(f'Активный классификатор не найден: {name}')
    return path

def clean(value):
    return re.sub(r'\s+', ' ', str(value)).strip() if value is not None else ''

def service_code(col):
    if 14 <= col <= 15: return '101'
    if 16 <= col <= 19: return 'PSC'
    if col == 20: return 'MGPSS'
    if 21 <= col <= 23: return '102'
    if 24 <= col <= 26: return '103'
    if 27 <= col <= 28: return '104'
    if 29 <= col <= 33: return 'CEMP'
    if 34 <= col <= 35: return 'FSB'
    special = {36:'MOSOBLGAZ',41:'GORHOZ',42:'GORMOST',43:'GORMOST',44:'GORMOST',45:'GORMOST',
               48:'EKP47',57:'MOSBEZ',58:'MOSBEZ_ANALYTICS',75:'TERRITORY',76:'TINAO',
               79:'EKP78',87:'EKP86',90:'EKP89',94:'EKP93',97:'EKP96'}
    return special.get(col, f'EKP{col}')

def parse_workbook(path):
    path = Path(path)
    workbook = load_workbook(path, data_only=True)
    sheet = workbook.active
    if sheet.max_column not in (90,99) or clean(sheet.cell(2,11).value) != 'Итоговый тип происшествия':
        raise ValueError('Структура XLSX не соответствует поддерживаемым версиям ЕКП (90/99 колонок).')
    merged = {}
    for area in sheet.merged_cells.ranges:
        if area.min_row <= 3:
            for row in range(area.min_row, min(area.max_row,3)+1):
                for col in range(area.min_col, area.max_col+1):
                    merged[row,col] = sheet.cell(area.min_row,area.min_col).value
    def header(row,col): return clean(merged.get((row,col), sheet.cell(row,col).value))
    columns = []
    legacy90=sheet.max_column==90
    if legacy90 and clean(sheet.cell(1,13).value)!='Сценарий реагирования':raise ValueError('Неподдерживаемая структура 90-колоночного ЕКП')
    for col in range(15 if legacy90 else 14,sheet.max_column+1):
        canonical=col-1 if legacy90 and col<=87 else col
        code = service_code(canonical)
        name = SERVICE_NAMES.get(code, header(1,col))
        columns.append({'column':col,'code':code,'name':name,'headers':[header(r,col) for r in (1,2,3)],
                        'condition':CONDITIONS.get(canonical, ('always', [])), 'visible':canonical!=58})
    section_names = {}
    for row in sheet.iter_rows(min_row=4,values_only=True):
        if row[4] is not None and not row[10] and row[5]: section_names[int(row[4])] = clean(row[5])
    # Group 24 has no standalone section row in the source workbook.
    section_names.setdefault(24, 'БПЛА')
    items, seen = [], set()
    for number,row in enumerate(sheet.iter_rows(min_row=4,values_only=True),4):
        if row[4] is None or not clean(row[10]): continue
        code = str(row[4])
        if code in seen: raise ValueError(f'Повтор кода ЕКП {code}, строка {number}')
        seen.add(code)
        rules = [{**column,'value':clean(row[column['column']-1])} for column in columns
                 if clean(row[column['column']-1])]
        items.append({'code':code,'group_code':str(row[0]),'category':section_names.get(row[0],f'Группа {row[0]}'),
                      'features':[clean(row[i]) for i in (6,7,8)], 'extra_questions':clean(row[9]),
                      'title':clean(row[10]),'legacy_title':clean(row[11]),'main_service':clean(row[13] if legacy90 else row[12]),'reaction_scenario':clean(row[12]) if legacy90 else '',
                      'source_row':number,'rules':rules,'raw':[v.isoformat() if hasattr(v,'isoformat') else v for v in row]})
    if not items: raise ValueError('В XLSX нет итоговых типов происшествия')
    workbook.close()
    return {'sha256':hashlib.sha256(path.read_bytes()).hexdigest(),'filename':path.name,'sheet':sheet.title,
            'groups':section_names,'columns':columns,'items':items}

def resolve_rules(rules,flags):
    """Resolve explicit header alternatives; preserve each contributing source column."""
    result = {}
    for rule in rules:
        mode, keys = rule['condition']
        selected = any(flags.get(key,False) for key in keys)
        applies = mode=='always' or (mode=='any' and selected) or (mode=='not' and not selected)
        if mode=='victims_here': applies = flags.get('victims',False) and not flags.get('victims_away',False)
        value = rule['value']
        if not applies or clean(value).lower() in ('нет реагирования','не реагирует','нет','-','0'): continue
        code = rule['code']
        entry = result.setdefault(code,{'code':code,'name':rule['name'],'visible':rule['visible'],'origin':'ЕКП','mappings':[]})
        entry['mappings'].append({'column':rule['column'],'type':value,'condition':rule['headers'][-1] or rule['headers'][1]})
    return list(result.values())
